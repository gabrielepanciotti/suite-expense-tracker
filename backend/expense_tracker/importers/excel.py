# expense_tracker/importers/excel.py
"""Parser del workbook storico delle spese (legacy .xlsx).

Struttura di un foglio spese:
    <CATEGORIA>                              (solo colonna A)
    Data | Descrizione | Importo | Fonte | ...   (riga intestazione)
    <righe dato...>
    TOTALE | | <somma> | ...

Una categoria è la riga "solo colonna A" che precede immediatamente la riga
intestazione `Data` (euristica robusta: gli header categoria contengono spesso
minuscole tra parentesi, es. "SPESA FISSA MENSILE (100 EUR/mese)", quindi NON ci
si può basare sul tutto-maiuscolo). Le righe nota/titolo (solo colonna A ma NON
seguite da un'intestazione `Data`) non diventano categorie.
"""
from __future__ import annotations

import logging
import os
from datetime import date, datetime

import openpyxl

from expense_tracker.domain import Expense, Kind, Payer, Source
from expense_tracker.importers.parsing import parse_amount_cents, parse_date

_log = logging.getLogger(__name__)

SHARED_SHEET = "Spese Condivise"
PERSONAL_SHEET = "Spese Personali"


def _sheet_candidates(env_var: str, default: str) -> list[str]:
    """Sheet-name candidates, env override first (call-time, device-local).

    The real legacy workbook may use a personalized sheet name: override via
    the untracked .env (LEGACY_SHARED_SHEET / LEGACY_PERSONAL_SHEET) — never
    hardcode personal names here. The generic default is always a fallback so
    sanitized fixtures parse regardless of the device env.
    """
    override = os.environ.get(env_var, "").strip()
    return [n for n in (override, default) if n]

_SOURCE_MAP = {
    "banca": Source.BANCA,
    "paypal": Source.PAYPAL,
    "manuale": Source.MANUALE,
    "screenshot": Source.SCREENSHOT,
    "contanti": Source.CONTANTI,
}


def _to_source(value: object) -> Source:
    return _SOURCE_MAP.get(str(value or "").strip().lower(), Source.BANCA)


def _coerce_date(value: object) -> date | None:
    """Cella data: datetime/date reale → usata direttamente; stringa → parse_date."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    return parse_date(str(value))


def _parse_sheet(ws, kind: Kind, shared: bool) -> list[Expense]:
    out: list[Expense] = []
    last_single: str | None = None
    category: str | None = None

    for row in ws.iter_rows(values_only=True):
        cells = tuple(row)
        a = cells[0]
        rest_empty = all(c in (None, "") for c in cells[1:])

        if rest_empty:
            # riga con sola colonna A (o del tutto vuota)
            if a not in (None, ""):
                last_single = str(a).strip()
            continue

        if str(a).strip() == "Data":
            category = last_single  # l'header categoria precede la riga 'Data'
            continue

        d = _coerce_date(a)
        if d is None:
            continue  # righe TOTALE/riepilogo/note con valori in altre colonne

        description = str(cells[1]).strip() if len(cells) > 1 and cells[1] not in (None, "") else ""
        if not description:
            continue
        try:
            cents = parse_amount_cents(cells[2])
        except (ValueError, IndexError) as e:
            _log.warning("riga ignorata (importo non parsabile): %s | %s | %r — %s",
                         d, description, cells[2] if len(cells) > 2 else None, e)
            continue

        out.append(
            Expense(
                date=d,
                description=description,
                amount_cents=cents,
                source=_to_source(cells[3] if len(cells) > 3 else None),
                category=category,
                kind=kind,
                shared=shared,
                payer=Payer.MEMBER_A,
            )
        )
    return out


def parse_workbook(path: str) -> list[Expense]:
    """Legge i fogli spese (condivise + personali) → lista di Expense."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: list[Expense] = []
    for name in _sheet_candidates("LEGACY_SHARED_SHEET", SHARED_SHEET):
        if name in wb.sheetnames:
            rows += _parse_sheet(wb[name], Kind.SPESA_CONDIVISA, shared=True)
            break
    for name in _sheet_candidates("LEGACY_PERSONAL_SHEET", PERSONAL_SHEET):
        if name in wb.sheetnames:
            rows += _parse_sheet(wb[name], Kind.SPESA_PERSONALE, shared=False)
            break

    # Indice di occorrenza fra righe altrimenti identiche (es. la stessa spesa
    # Decathlon listata sia in una categoria condivisa sia in una personale):
    # le preserva entrambe e mantiene l'idempotenza del re-import.
    seen: dict[tuple, int] = {}
    for r in rows:
        key = (r.date, r.description, r.amount_cents, r.source.value)
        r.seq = seen.get(key, 0)
        seen[key] = r.seq + 1
    return rows
