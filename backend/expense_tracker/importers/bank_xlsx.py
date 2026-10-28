"""Importer dell'estratto conto BBVA (xlsx).

Foglio unico `Informe BBVA`. Intestazione su riga 5; colonne (1-indexed):
A=vuota, B=`Data valuta`, C=`Data`, D=`Causale`, E=`Movimento`,
F=`Beneficiario`, G=`Importo`. Le righe dato iniziano a riga 6 e sono valide
solo se la colonna G (Importo) è non vuota.

Convenzione segno `amount_cents` (coerente con lo storico Excel: le spese sono
positive; `report` separa per kind):
- spesa_* / entrata / investimento → `abs(amount)` (magnitudine sempre positiva)
- member_transfer / giroconto → importo bancario *signed* così com'è (negativo =
  uscita/inviato, positivo = entrata/ricevuto): la direzione serve al saldo.

Counterparty matching is env-driven (NO personal names in tracked files):
- BANK_SELF_NAMES — ";"-separated names identifying the account holder
  (self-transfers → giroconto). Empty = rule inactive.
- BANK_MEMBER_B_NAMES — ";"-separated names identifying the other household
  member (transfers → member_transfer). Empty = rule inactive.
Set them in the untracked .env (see .env.example for fictional samples).

Tutte le righe sono provvisorie: la successiva `review` le finalizza. L'importer
non assegna alcuno status (è compito del chiamante).
"""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

from expense_tracker.domain import Expense, Kind, Payer, Source
from expense_tracker.importers.parsing import (
    collapse_ws,
    env_names,
    parse_bbva_amount,
    parse_date,
)

_log = logging.getLogger(__name__)

_FIRST_DATA_ROW = 6

# colonne (1-indexed)
_COL_DATA_VALUTA = 2  # B
_COL_CAUSALE = 4  # D
_COL_MOVIMENTO = 5  # E
_COL_BENEFICIARIO = 6  # F
_COL_IMPORTO = 7  # G

_SHEET = "Informe BBVA"

# Movements identifying an internal transfer between the user's own accounts.
_GIROCONTO_MOVIMENTI = {
    "Inviato da Revolut",
    "INSTANT TRANSFER",
    "Trasferimento su altro conto",
}
_BONIFICO_CAUSALI = ("BONIFICO ESEGUITO", "BONIFICO RICEVUTO")


def _first_line(value: object) -> str:
    """Prima riga del Beneficiario; tratta '-' come vuoto."""
    first = str(value or "").splitlines()[0].strip() if value not in (None, "") else ""
    return "" if first == "-" else first


def _classify(causale: str, movimento: str, benef: str, amount: int) -> dict:
    """Inferisce description/kind/category/amount applicando le regole in ordine
    (prima corrispondenza vince). Ritorna i campi di classificazione (senza
    date/source/payer): il chiamante costruisce la Expense finale.
    """
    abs_amount = abs(amount)

    # 1. Pagamento con carta → spesa personale (merchant = causale)
    if movimento == "PAGAMENTO CON CARTA":
        return dict(description=causale, amount_cents=abs_amount,
                    kind=Kind.SPESA_PERSONALE, category=None)

    # 2. Stipendio
    if causale.startswith("ACCREDITO DELLO STIPENDIO"):
        return dict(description=f"Stipendio {movimento}".strip(),
                    amount_cents=abs_amount, kind=Kind.ENTRATA, category="STIPENDIO")

    # 3. Interessi/competenze bancarie
    if causale.startswith("LIQUIDAZIONE INTERESSI"):
        return dict(description="Interessi/competenze bancarie",
                    amount_cents=abs_amount, kind=Kind.ENTRATA,
                    category="INTERESSI E COMPETENZE")

    # 4. PayPal (addebito cumulativo, dettaglio a posteriori)
    if "PayPal" in causale:
        return dict(description="PayPal (da dettagliare)", amount_cents=abs_amount,
                    kind=Kind.SPESA_PERSONALE, category=None)

    # 5. Scalable Capital (PAC investimento)
    if "Scalable Capital" in causale:
        return dict(description="Scalable Capital (PAC)", amount_cents=abs_amount,
                    kind=Kind.INVESTIMENTO, category="PAC Scalable")

    # 6. Other household member (counterparty in BANK_MEMBER_B_NAMES) →
    #    member_transfer, importo signed
    benef_cf = benef.casefold()
    if any(name in benef_cf for name in env_names("BANK_MEMBER_B_NAMES")):
        return dict(description=f"{movimento} — Member B".strip(),
                    amount_cents=amount, kind=Kind.MEMBER_TRANSFER, category=None)

    # 7. Bonifico eseguito/ricevuto
    if causale in _BONIFICO_CAUSALI:
        is_giroconto = (
            movimento in _GIROCONTO_MOVIMENTI
            or "Scalable Capital" in movimento
            or (
                causale == "BONIFICO ESEGUITO"
                and benef.casefold() in env_names("BANK_SELF_NAMES")
            )
        )
        label = movimento or causale
        desc = collapse_ws(f"{label} — {benef}" if benef else label)
        if is_giroconto:
            return dict(description=desc, amount_cents=amount,
                        kind=Kind.GIROCONTO, category="GIROCONTO")
        # bonifico verso/da terzi reale
        kind = Kind.SPESA_PERSONALE if amount < 0 else Kind.ENTRATA
        return dict(description=desc, amount_cents=abs_amount, kind=kind, category=None)

    # 8. Fallback
    kind = Kind.SPESA_PERSONALE if amount < 0 else Kind.ENTRATA
    return dict(description=causale, amount_cents=abs_amount, kind=kind, category=None)


def parse_bbva(path: str | Path) -> list[Expense]:
    """Legge l'estratto conto BBVA → lista di Expense provvisorie."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[_SHEET] if _SHEET in wb.sheetnames else wb.active

    out: list[Expense] = []
    for r in range(_FIRST_DATA_ROW, ws.max_row + 1):
        importo_raw = ws.cell(row=r, column=_COL_IMPORTO).value
        if importo_raw in (None, ""):
            continue  # riga non-dato (footer/vuota)

        try:
            amount = parse_bbva_amount(importo_raw)
        except (ValueError, TypeError) as e:
            _log.warning("riga %d ignorata (importo non parsabile): %r — %s", r, importo_raw, e)
            continue

        d = parse_date(str(ws.cell(row=r, column=_COL_DATA_VALUTA).value or ""))
        causale = collapse_ws(ws.cell(row=r, column=_COL_CAUSALE).value)
        movimento = collapse_ws(ws.cell(row=r, column=_COL_MOVIMENTO).value)
        benef = _first_line(ws.cell(row=r, column=_COL_BENEFICIARIO).value)

        fields = _classify(causale, movimento, benef, amount)
        out.append(Expense(
            date=d, source=Source.BANCA, payer=Payer.MEMBER_A,
            shared=None,
            description=collapse_ws(fields["description"]),
            amount_cents=fields["amount_cents"],
            kind=fields["kind"], category=fields["category"],
        ))

    # Indice di occorrenza fra righe altrimenti identiche (es. due addebiti PayPal
    # nello stesso giorno e stesso importo → entrambe "PayPal (da dettagliare)"):
    # senza `seq` condividerebbero lo stesso raw_hash e la seconda verrebbe persa.
    seen: dict[tuple, int] = {}
    for t in out:
        key = (t.date, t.description, t.amount_cents)
        t.seq = seen.get(key, 0)
        seen[key] = t.seq + 1
    return out
